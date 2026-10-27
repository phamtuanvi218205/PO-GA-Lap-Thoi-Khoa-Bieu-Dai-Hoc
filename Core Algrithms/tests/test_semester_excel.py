"""Kiểm thử workbook nhập dữ liệu học kỳ."""

from io import BytesIO

import pytest
from openpyxl import load_workbook

from gapo_timetabling.application.semester_excel import (
    SemesterExcelService,
    SemesterExcelValidationError,
)


def test_generated_template_can_be_parsed_by_the_same_contract() -> None:
    service = SemesterExcelService()

    content = service.create_template()
    bundle = service.parse(content)

    assert content.startswith(b"PK")
    assert bundle.term.term_code == "2026_HK1"
    assert bundle.term.number_of_weeks == 15
    workbook = load_workbook(BytesIO(content), read_only=True)
    assert "CauTrucMonHoc" in workbook.sheetnames
    assert "ThietBiCauTruc" in workbook.sheetnames
    assert "PhanGiangDay" not in workbook.sheetnames


def test_parser_reports_sheet_row_and_field_for_invalid_value() -> None:
    service = SemesterExcelService()
    workbook = load_workbook(BytesIO(service.create_template()))
    workbook["HocKy"]["D2"] = "không phải số"
    output = BytesIO()
    workbook.save(output)

    with pytest.raises(SemesterExcelValidationError) as captured:
        service.parse(output.getvalue())

    issue = captured.value.issues[0]
    assert issue.sheet == "HocKy"
    assert issue.row == 2
    assert issue.field == "number_of_weeks"


def test_parser_rejects_modified_headers_instead_of_guessing_columns() -> None:
    service = SemesterExcelService()
    workbook = load_workbook(BytesIO(service.create_template()))
    workbook["MonHoc"]["A1"] = "Mã môn tự đặt"
    output = BytesIO()
    workbook.save(output)

    with pytest.raises(SemesterExcelValidationError, match="Workbook có") as captured:
        service.parse(output.getvalue())

    assert any(issue.sheet == "MonHoc" and issue.row == 1 for issue in captured.value.issues)
