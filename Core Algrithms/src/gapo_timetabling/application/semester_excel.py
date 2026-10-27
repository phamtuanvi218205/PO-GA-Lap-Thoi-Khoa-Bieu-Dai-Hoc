"""Đọc và tạo workbook nhập dữ liệu học kỳ.

Workbook chỉ là một giao diện nhập hàng loạt.  Sau khi đọc, mọi dòng được
chuyển thành ``SemesterInputBundle`` và đi qua cùng validator/service với API
nhập tay.  Importer không tự đoán mã tham chiếu hoặc tự sửa dữ liệu sai.
"""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from typing import Any, Callable

from pydantic import BaseModel, ValidationError

from .input_data import (
    AcademicTermInput,
    CampusInput,
    CourseInput,
    CourseSectionInput,
    CourseTeachingStructureEquipmentInput,
    CourseTeachingStructureInput,
    EquipmentInput,
    LecturerAvailabilityInput,
    LecturerCapabilityInput,
    LecturerInput,
    OnlineLocationInput,
    RoomAvailabilityInput,
    RoomEquipmentInput,
    RoomInput,
    SemesterInputBundle,
)


@dataclass(frozen=True, slots=True)
class ExcelImportIssue:
    sheet: str
    row: int | None
    field: str | None
    message: str


class SemesterExcelValidationError(ValueError):
    """Workbook sai cấu trúc hoặc chứa dòng không hợp lệ."""

    def __init__(self, issues: list[ExcelImportIssue]) -> None:
        super().__init__(f"Workbook có {len(issues)} lỗi dữ liệu.")
        self.issues = tuple(issues)


@dataclass(frozen=True, slots=True)
class _SheetSpec:
    name: str
    headers: tuple[str, ...]
    model: type[BaseModel]
    target_field: str
    converters: dict[str, Callable[[Any], Any]]
    required: bool = False


def _text(value: Any) -> str | None:
    if value is None:
        return None
    result = str(value).strip()
    return result or None


def _integer(value: Any) -> int | None:
    if value is None or str(value).strip() == "":
        return None
    if isinstance(value, bool):
        raise ValueError("Giá trị số nguyên không được là boolean.")
    number = float(value)
    if not number.is_integer():
        raise ValueError("Giá trị phải là số nguyên.")
    return int(number)


def _decimal(value: Any) -> Decimal | None:
    if value is None or str(value).strip() == "":
        return None
    return Decimal(str(value).replace(",", "."))


def _boolean(value: Any) -> bool | None:
    if value is None or str(value).strip() == "":
        return None
    if isinstance(value, bool):
        return value
    normalized = str(value).strip().lower()
    if normalized in {"1", "true", "yes", "y", "có", "co", "x"}:
        return True
    if normalized in {"0", "false", "no", "n", "không", "khong"}:
        return False
    raise ValueError("Giá trị boolean phải là Có/Không hoặc True/False.")


def _date(value: Any) -> date | None:
    if value is None or str(value).strip() == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    for pattern in ("%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(text, pattern).date()
        except ValueError:
            pass
    raise ValueError("Ngày phải có dạng YYYY-MM-DD hoặc DD/MM/YYYY.")


_SPECS: tuple[_SheetSpec, ...] = (
    _SheetSpec(
        "HocKy",
        ("term_code", "term_name", "start_date", "number_of_weeks", "status"),
        AcademicTermInput,
        "term",
        {"start_date": _date, "number_of_weeks": _integer},
        required=True,
    ),
    _SheetSpec(
        "CoSo",
        ("campus_code", "campus_name", "address", "is_active"),
        CampusInput,
        "campuses",
        {"is_active": _boolean},
    ),
    _SheetSpec(
        "MonHoc",
        ("course_code", "course_name", "credits", "is_active"),
        CourseInput,
        "courses",
        {"credits": _decimal, "is_active": _boolean},
    ),
    _SheetSpec(
        "GiangVien",
        (
            "lecturer_code", "full_name", "email", "home_campus_code",
            "max_teaching_periods", "is_active",
        ),
        LecturerInput,
        "lecturers",
        {"max_teaching_periods": _integer, "is_active": _boolean},
    ),
    _SheetSpec(
        "NangLucGiangDay",
        (
            "lecturer_code", "course_code", "can_teach_lecture",
            "can_teach_practice", "suitability_score",
        ),
        LecturerCapabilityInput,
        "lecturer_capabilities",
        {
            "can_teach_lecture": _boolean,
            "can_teach_practice": _boolean,
            "suitability_score": _decimal,
        },
    ),
    _SheetSpec(
        "PhongHoc",
        ("room_code", "room_name", "campus_code", "room_type", "capacity", "is_active"),
        RoomInput,
        "rooms",
        {"capacity": _integer, "is_active": _boolean},
    ),
    _SheetSpec(
        "LopTrucTuyen",
        ("location_code", "location_name", "platform_name", "meeting_note", "is_active"),
        OnlineLocationInput,
        "online_locations",
        {"is_active": _boolean},
    ),
    _SheetSpec(
        "ThietBi",
        ("equipment_code", "equipment_name"),
        EquipmentInput,
        "equipment",
        {},
    ),
    _SheetSpec(
        "ThietBiPhong",
        ("room_code", "equipment_code", "quantity"),
        RoomEquipmentInput,
        "room_equipment",
        {"quantity": _integer},
    ),
    _SheetSpec(
        "LopHocPhan",
        ("section_code", "course_code", "section_name", "expected_enrollment", "status"),
        CourseSectionInput,
        "course_sections",
        {"expected_enrollment": _integer},
    ),
    _SheetSpec(
        "CauTrucMonHoc",
        (
            "course_code", "part_type", "total_periods",
            "required_location_type", "required_room_type",
            "max_sessions_per_week", "allows_intensive", "max_candidate_plans",
            "is_active", "note",
        ),
        CourseTeachingStructureInput,
        "course_teaching_structures",
        {
            "total_periods": _integer,
            "max_sessions_per_week": _integer,
            "allows_intensive": _boolean,
            "max_candidate_plans": _integer,
            "is_active": _boolean,
        },
    ),
    _SheetSpec(
        "ThietBiCauTruc",
        ("course_code", "part_type", "equipment_code", "minimum_quantity"),
        CourseTeachingStructureEquipmentInput,
        "course_structure_equipment",
        {"minimum_quantity": _integer},
    ),
    _SheetSpec(
        "LichGiangVien",
        (
            "lecturer_code", "scope_type", "calendar_date", "iso_weekday",
            "start_period", "end_period", "availability_type",
            "preference_weight", "note",
        ),
        LecturerAvailabilityInput,
        "lecturer_availability",
        {
            "calendar_date": _date,
            "iso_weekday": _integer,
            "start_period": _integer,
            "end_period": _integer,
            "preference_weight": _decimal,
        },
    ),
    _SheetSpec(
        "LichPhong",
        (
            "room_code", "scope_type", "calendar_date", "iso_weekday",
            "start_period", "end_period", "availability_type", "note",
        ),
        RoomAvailabilityInput,
        "room_availability",
        {
            "calendar_date": _date,
            "iso_weekday": _integer,
            "start_period": _integer,
            "end_period": _integer,
        },
    ),
)


class SemesterExcelService:
    """Chuyển đổi hai chiều giữa workbook mẫu và hợp đồng dữ liệu chuẩn."""

    def parse(self, content: bytes) -> SemesterInputBundle:
        if not content:
            raise SemesterExcelValidationError([
                ExcelImportIssue("Workbook", None, None, "Tệp Excel rỗng.")
            ])
        try:
            from openpyxl import load_workbook

            workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        except Exception as error:
            raise SemesterExcelValidationError([
                ExcelImportIssue("Workbook", None, None, f"Không đọc được tệp .xlsx: {error}")
            ]) from error

        issues: list[ExcelImportIssue] = []
        parsed: dict[str, Any] = {}
        for spec in _SPECS:
            if spec.name not in workbook.sheetnames:
                if spec.required:
                    issues.append(ExcelImportIssue(spec.name, None, None, "Thiếu worksheet bắt buộc."))
                else:
                    parsed[spec.target_field] = []
                continue
            sheet = workbook[spec.name]
            rows = sheet.iter_rows(values_only=True)
            header_row = next(rows, None)
            headers = tuple(_text(value) or "" for value in (header_row or ()))
            if headers != spec.headers:
                issues.append(ExcelImportIssue(
                    spec.name,
                    1,
                    None,
                    "Tiêu đề cột không đúng mẫu. Cần: " + ", ".join(spec.headers),
                ))
                parsed[spec.target_field] = []
                continue
            records: list[BaseModel] = []
            for row_number, values in enumerate(rows, start=2):
                if all(value is None or str(value).strip() == "" for value in values):
                    continue
                raw = dict(zip(spec.headers, values, strict=False))
                converted: dict[str, Any] = {}
                row_failed = False
                for field, value in raw.items():
                    converter = spec.converters.get(field, _text)
                    try:
                        converted[field] = converter(value)
                    except (ValueError, TypeError, ArithmeticError) as error:
                        issues.append(ExcelImportIssue(spec.name, row_number, field, str(error)))
                        row_failed = True
                if row_failed:
                    continue
                converted = {
                    field: value
                    for field, value in converted.items()
                    if value is not None
                    or spec.model.model_fields[field].is_required()
                }
                try:
                    records.append(spec.model.model_validate(converted))
                except ValidationError as error:
                    for detail in error.errors():
                        field = ".".join(map(str, detail["loc"])) or None
                        issues.append(ExcelImportIssue(
                            spec.name, row_number, field, str(detail["msg"])
                        ))
            if spec.target_field == "term":
                if len(records) != 1:
                    issues.append(ExcelImportIssue(
                        spec.name, None, None, "Sheet HocKy phải có đúng một dòng dữ liệu."
                    ))
                elif records:
                    parsed[spec.target_field] = records[0]
            else:
                parsed[spec.target_field] = records

        if issues:
            raise SemesterExcelValidationError(issues)
        try:
            return SemesterInputBundle.model_validate(parsed)
        except ValidationError as error:
            bundle_issues = [
                ExcelImportIssue(
                    "Workbook",
                    None,
                    ".".join(map(str, detail["loc"])) or None,
                    str(detail["msg"]),
                )
                for detail in error.errors()
            ]
            raise SemesterExcelValidationError(bundle_issues) from error

    def create_template(self) -> bytes:
        """Tạo workbook trống với tiêu đề, hướng dẫn ngắn và định dạng nhập liệu."""

        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.worksheet.datavalidation import DataValidation

        workbook = Workbook()
        workbook.remove(workbook.active)
        header_fill = PatternFill("solid", fgColor="0F766E")
        guide = workbook.create_sheet("HuongDan")
        guide.append(["MẪU NHẬP DỮ LIỆU XẾP THỜI KHÓA BIỂU"])
        guide.append(["Mỗi mã phải duy nhất và phải khớp chính xác giữa các sheet."])
        guide.append(["Chỉ nhập dữ liệu biết trước khi xếp; không nhập TeachingPlan, session, option hoặc lịch kết quả."])
        guide.append(["Ngày dùng YYYY-MM-DD hoặc DD/MM/YYYY. Thứ dùng ISO: 1 = Thứ Hai, 7 = Chủ nhật."])
        guide.append(["Khai báo cấu trúc LECTURE/PRACTICE một lần cho mỗi môn; hệ thống tự sinh phần giảng dạy cho các lớp được mở."])
        guide.append(["required_location_type: PHYSICAL_ROOM/ONLINE."])
        guide.append(["availability_type giảng viên: UNAVAILABLE/DISCOURAGED/PREFERRED."])
        guide.append(["Hãy gọi API preview và sửa hết lỗi trước khi áp dụng vào database."])
        guide["A1"].font = Font(bold=True, size=14, color="17324D")
        guide.column_dimensions["A"].width = 115
        for spec in _SPECS:
            sheet = workbook.create_sheet(spec.name)
            sheet.append(spec.headers)
            for cell in sheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = f"A1:{_column_letter(len(spec.headers))}1"
            for column_index, header in enumerate(spec.headers, start=1):
                sheet.column_dimensions[_column_letter(column_index)].width = max(14, min(28, len(header) + 4))

        term = workbook["HocKy"]
        term.append(["2026_HK1", "Học kỳ 1 năm học 2026-2027", date(2026, 8, 17), 15, "PLANNING"])
        term["C2"].number_format = "dd/mm/yyyy"

        for sheet_name, columns in {
            "CoSo": (4,), "MonHoc": (4,), "GiangVien": (6,),
            "NangLucGiangDay": (3, 4), "PhongHoc": (6,),
            "LopTrucTuyen": (5,), "CauTrucMonHoc": (7, 9),
        }.items():
            sheet = workbook[sheet_name]
            for column in columns:
                boolean_validation = DataValidation(
                    type="list", formula1='"Có,Không"', allow_blank=True
                )
                sheet.add_data_validation(boolean_validation)
                boolean_validation.add(f"{_column_letter(column)}2:{_column_letter(column)}1000")

        output = BytesIO()
        workbook.save(output)
        return output.getvalue()


def _column_letter(index: int) -> str:
    result = ""
    while index:
        index, remainder = divmod(index - 1, 26)
        result = chr(65 + remainder) + result
    return result
